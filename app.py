# =============STEP 1: LOAD MODULES===============
import base64
import pandas as pd
import streamlit as st
from docx import Document as DocxDocument
from pypdf import PdfReader
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_core.prompts import PromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnablePassthrough
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_google_genai import ChatGoogleGenerativeAI

# Change the Gemini model name here ONCE, it is used everywhere below
MODEL_NAME = "gemini-3.5-flash"

# ============= STEP 2: PAGE CONFIGURATION ===============

st.set_page_config(page_title="Chat with Medical Reports", page_icon="🩺", layout="wide")
st.title("🩺 Chat with Medical Reports")
st.caption("Upload a lab report or prescription and ask questions in plain English.")

st.markdown(
    """
    <style>
    [data-testid="stSidebar"] {
        background-color: #001f3f;
    }
    /* Make sidebar text white so it stays readable on navy blue */
    [data-testid="stSidebar"] * {
        color: #ffffff;
    }
    /* Style the text input box so it doesn't stay white-on-white */
    [data-testid="stSidebar"] input {
        background-color: #0a2d5c;
        color: #ffffff;
    }
    [data-testid="stFileUploaderDropzone"] { background-color: #0a2d5c; }
    [data-testid="stSidebar"] hr {
        border-color: #ffffff;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ============ STEP 3: API KEYS & FILE UPLOAD =================

with st.sidebar:
    st.header("USER DETAILS")
    st.divider()
    st.header("Enter your API Key: ")
    GOOGLE_API_KEY = st.text_input("google_api_key", type="password",
                                   help="Get a free key at aistudio.google.com/app/apikey")
    st.header("Upload your file")
    uploaded_file = st.file_uploader(
        "Upload medical report",
        type=["pdf", "docx", "txt", "csv", "xlsx", "png", "jpg", "jpeg"],
    )
    st.divider()
    st.markdown(
        """⚠️ **Disclaimer:** This tool explains report content in plain 
        language for educational purposes only. It is **not** a medical 
        diagnosis. Always consult a qualified doctor."""
    )

# ============ HEALTH SYSTEM PROMPTS ==============

MEDICAL_PROMPT = """You are a careful medical-report explainer assistant.
You help patients understand their lab reports and prescriptions in simple,
plain language. You are NOT a doctor and must never provide a diagnosis or
tell the user to start/stop/change any medication or treatment.

Rules you must always follow:
1. Base your answer ONLY on the provided report context. If the answer isn't
   in the context, SAY "I don't have such kind of knowledge".
2. When explaining a lab value, mention: what the test measures, the
   patient's value (if present), the normal/reference range (if present),
   and whether it appears high, low, or normal — in plain, non-technical
   language a layperson can understand.
3. Avoid medical jargon; if you must use a technical term, explain it in
   parentheses.
4. Never give a diagnosis, prognosis, or specific treatment/dosage advice.
   Instead, gently suggest the person discuss the finding with their doctor.
5. Keep tone calm and reassuring — avoid alarming language even for
   abnormal values.
6. End every answer with a short one-line reminder to consult a healthcare
   professional, EXCEPT for simple factual lookups (e.g. "what is my
   hemoglobin value") where it isn't needed.

Context from the uploaded report:
{context}"""

QA_TEMPLATE = MEDICAL_PROMPT + """

Question: {question}

Answer:"""

QA_PROMPT = PromptTemplate(
    template=QA_TEMPLATE,
    input_variables=["context", "question"],
)

SUMMARY_TEMPLATE = MEDICAL_PROMPT + """

Give a short, plain-language summary of this report: what tests were done,
which values (if any) are outside the normal range, and what that might
generally mean. Keep it under 200 words.

Answer:"""

SUMMARY_PROMPT = PromptTemplate(
    template=SUMMARY_TEMPLATE,
    input_variables=["context"],
)

# Used to suggest which type of doctor to consult (shown on main page + Find Doctors page)
GUIDE_TEMPLATE = """Based on this medical report, suggest 1 to 3 types of doctors/specialists
the person should consider consulting (e.g. General Physician, Endocrinologist) and
give a one-line reason for each. Do NOT diagnose or name diseases as certain.

Report content:
{context}

Suggestions:"""


def format_docs(docs):
    return "\n\n".join(d.page_content for d in docs)

# =============== READ DIFFERENT FILE TYPES =====================

def extract_pages(uploaded_file, api_key):
    ext = uploaded_file.name.lower().split(".")[-1]

    if ext == "pdf":
        reader = PdfReader(uploaded_file)
        texts = [p.extract_text() or "" for p in reader.pages]

    elif ext == "docx":
        doc = DocxDocument(uploaded_file)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:                      # lab reports are often tables
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        texts = ["\n".join(lines)]

    elif ext == "txt":
        texts = [uploaded_file.read().decode("utf-8", errors="ignore")]

    elif ext == "csv":
        texts = [pd.read_csv(uploaded_file).to_string(index=False)]

    elif ext == "xlsx":
        sheets = pd.read_excel(uploaded_file, sheet_name=None)
        texts = [f"Sheet: {name}\n{df.to_string(index=False)}" for name, df in sheets.items()]

    elif ext in ("png", "jpg", "jpeg"):
        # Use Gemini's vision to read the image, so no Tesseract install is needed
        img_b64 = base64.b64encode(uploaded_file.read()).decode()
        mime = "jpeg" if ext in ("jpg", "jpeg") else "png"
        llm = ChatGoogleGenerativeAI(model=MODEL_NAME, google_api_key=api_key)
        msg = HumanMessage(content=[
            {"type": "text", "text": "Extract all text from this medical document exactly as written, including tables and values."},
            {"type": "image_url", "image_url": f"data:image/{mime};base64,{img_b64}"},
        ])
        texts = [llm.invoke([msg]).content]

    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    return [Document(page_content=t, metadata={"page": i}) for i, t in enumerate(texts)]

# =============== RAG PIPELINE =====================

@st.cache_resource(show_spinner=False)
def build_vectorstore(uploaded_file, api_key):
    pages = extract_pages(uploaded_file, api_key)

    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=150)
    chunks = splitter.split_documents(pages)

    embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
    return FAISS.from_documents(chunks, embeddings)


def build_chain(vectorstore, api_key, prompt_template):
    retriever = vectorstore.as_retriever(search_kwargs={"k": 4})
    llm = ChatGoogleGenerativeAI(model=MODEL_NAME, google_api_key=api_key)

    rag_chain = (
        {"context": retriever | format_docs, "question": RunnablePassthrough()}
        | prompt_template
        | llm
        | StrOutputParser()
    )
    return rag_chain

# ============ MAIN APP ===============

if "messages" not in st.session_state:
    st.session_state.messages = []

if uploaded_file and not GOOGLE_API_KEY:
    st.warning("Please enter your Google API key in the sidebar to continue.")

if uploaded_file and GOOGLE_API_KEY:
    with st.spinner("Reading and indexing your report..."):
        vectorstore = build_vectorstore(uploaded_file, GOOGLE_API_KEY)
    st.success(f"'{uploaded_file.name}' processed successfully. Ask away!")

    # Share these with the "Find Doctors" page
    st.session_state["vectorstore"] = vectorstore
    st.session_state["google_api_key"] = GOOGLE_API_KEY

    # Generate specialist suggestions once per uploaded file
    if st.session_state.get("guidance_for") != uploaded_file.name:
        st.session_state.messages = []   # new report -> start a fresh chat
        with st.spinner("Preparing guidance for your report..."):
            docs = vectorstore.as_retriever(search_kwargs={"k": 6}).invoke(
                "abnormal values, diagnosis, medicines, doctor remarks")
            guide_chain = (
                PromptTemplate.from_template(GUIDE_TEMPLATE)
                | ChatGoogleGenerativeAI(model=MODEL_NAME, google_api_key=GOOGLE_API_KEY, temperature=0.2)
                | StrOutputParser()
            )
            st.session_state["guidance"] = guide_chain.invoke({"context": format_docs(docs)})
        st.session_state["guidance_for"] = uploaded_file.name

    st.info("🩺 **Suggested specialists**\n\n" + st.session_state["guidance"])
    st.success("👉 Open the **Find Doctors** page from the sidebar to see doctors and hospitals near you.")

    col1, col2 = st.columns([1, 1])
    with col1:
        if st.button("📋 Get a quick plain-language summary", use_container_width=True):
            summary_chain = build_chain(vectorstore, GOOGLE_API_KEY, SUMMARY_PROMPT)
            with st.spinner("Summarizing report..."):
                summary = summary_chain.invoke("Summarize this report.")
            st.session_state.messages.append({"role": "assistant", "content": summary})

    # Display chat history
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    # Chat input
    if question := st.chat_input("Ask a question about your report..."):
        st.session_state.messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)

        qa_chain = build_chain(vectorstore, GOOGLE_API_KEY, QA_PROMPT)
        with st.chat_message("assistant"):
            with st.spinner("Thinking..."):
                answer = qa_chain.invoke(question)
                st.markdown(answer)
        st.session_state.messages.append({"role": "assistant", "content": answer})

elif not uploaded_file:
    st.info("👈 Upload your medical report (PDF, Word, image, etc.) in the sidebar to get started.")
